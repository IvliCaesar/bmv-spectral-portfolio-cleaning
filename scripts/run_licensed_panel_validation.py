from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import minimize
from scipy.stats import norm
from sklearn.covariance import LedoitWolf

ROOT = Path(__file__).resolve().parents[1]


ESTIMATORS = ("Sample", "Eigenvalue clipping", "Ledoit-Wolf", "RIE")
TEST_FRACTION = 0.20
ROLLING_LOOKBACK = 504
ROLLING_STEP = 21
ROLLING_REBALANCES = 25
ANNUAL_RISK_FREE_RATE = 0.08
BOOTSTRAP_REPLICATIONS = 20_000
BOOTSTRAP_SEED = 20260930
START_DATE = pd.Timestamp("2011-07-18")
END_DATE = pd.Timestamp("2025-09-29")
TICKER_SHEET_MAP = {
    "MFRISCOA-1": "MFRISCOA MM Equity",
    "RA": "RA MM Equity",
    "SIMECB": "SIMECB MM Equity",
    "ICHB": "ICHB MM Equity",
    "GENTERA": "GENTERA  MM Equity",
    "SORIANAB": "SORIANAB MM Equity",
    "ELEKTRA": "ELEKTRA  MM Equity",
    "CHDRAUIB": "CHDRAUIB MM Equity",
    "LIVEPOLC-1": "LIVEPOLC MM Equity",
    "AUTLANB": "AUTLANB MM Equity",
    "HERDEZ": "HERDEZ  MM Equity",
    "GAPB": "GAPB MM Equity",
    "OMAB": "OMAB MM Equity",
    "MEGACPO": "MEGACPO MM Equity",
    "PINFRA": "PINFRA  MM Equity",
    "LABB": "LABB MM Equity",
    "ALSEA": "ALSEA  MM Equity",
    "ARA": "ARA  MM Equity",
    "ASURB": "ASURB MM Equity",
    "AMXB": "AMXB MM Equity",
    "AC": "AC  MM Equity",
    "AXTELCPO": "AXTELCPO MM Equity",
    "KIMBERA": "KIMBERA MM Equity",
    "CEMEXCPO": "CEMEXCPO MM Equity",
    "BIMBOA": "BIMBOA MM Equity",
    "BOLSAA": "BOLSAA MM Equity",
    "GCARSOA1": "GCARSOA1 MM Equity",
    "FEMSAUBD": "FEMSAUBD MM Equity",
    "GFNORTEO": "GFNORTEO MM Equity",
    "GFINBURO": "GFINBURO MM Equity",
    "ORBIA": "ORBIA  MM Equity",
    "GRUMAB": "GRUMAB MM Equity",
    "GMEXICOB": "GMEXICOB MM Equity",
    "PE&OLES": "PE&OLES  MM Equity",
    "TLEVISACPO": "TLEVICPO MM Equity",
    "SIGMAFA": "SIGMAFA MM Equity",
    "WALMEX": "WALMEX  MM Equity",
}


def validate_daily_dates(dates: pd.Series) -> None:
    gaps = dates.sort_values().diff().dt.days
    large_gaps = gaps[gaps > 7]
    if not large_gaps.empty:
        raise ValueError(
            "Input is not a continuous daily trading-date panel: "
            f"{len(large_gaps)} gaps exceed seven calendar days (maximum "
            f"{int(large_gaps.max())} days). Do not annualize or label "
            "observation-count windows as trading days; recover the missing "
            "daily total-return observations first."
        )


def gmv_portfolio(covariance: np.ndarray) -> np.ndarray:
    ones = np.ones(covariance.shape[0])
    raw = np.linalg.solve(covariance, ones)
    return raw / (ones @ raw)


def eigenvalue_clipping(
    covariance: np.ndarray, assets: int, observations: int
) -> np.ndarray:
    upper_edge = (1 + np.sqrt(assets / observations)) ** 2
    standard_deviations = np.sqrt(np.diag(covariance))
    correlation = covariance / np.outer(standard_deviations, standard_deviations)
    eigenvalues, eigenvectors = np.linalg.eigh(correlation)
    eigenvalues = eigenvalues[::-1]
    eigenvectors = eigenvectors[:, ::-1]
    signal_count = max(int(np.sum(eigenvalues > upper_edge)), 1)
    if signal_count >= assets:
        raise RuntimeError("Eigenvalue clipping found no noise eigenvalues to average.")
    noise_mean = eigenvalues[signal_count:].mean()
    clipped = np.where(np.arange(assets) < signal_count, eigenvalues, noise_mean)
    cleaned = eigenvectors @ np.diag(clipped) @ eigenvectors.T
    diagonal = np.diag(cleaned)
    cleaned /= np.outer(np.sqrt(diagonal), np.sqrt(diagonal))
    return np.outer(standard_deviations, standard_deviations) * cleaned


def rie_estimator(covariance: np.ndarray, assets: int, observations: int) -> np.ndarray:
    aspect_ratio = assets / observations
    eigenvalues, eigenvectors = np.linalg.eigh(covariance)
    bandwidth = eigenvalues * observations ** (-1.0 / 3.0)
    difference = eigenvalues[:, None] - eigenvalues[None, :]
    relative_difference = difference / bandwidth[None, :]
    bracket = 1.0 - relative_difference**2 / 5.0
    positive_bracket = np.clip(bracket, 0.0, None)
    density = (
        3.0
        / (4 * np.sqrt(5) * bandwidth[None, :])
        * positive_bracket
    ).mean(axis=1)
    term_one = -3.0 * difference / (10 * np.pi * bandwidth[None, :] ** 2)
    numerator = np.sqrt(5) * bandwidth[None, :] - difference
    denominator = np.sqrt(5) * bandwidth[None, :] + difference
    with np.errstate(divide="ignore", invalid="ignore"):
        logarithm = np.log(np.abs(numerator / denominator))
    logarithm = np.where(np.isfinite(logarithm), logarithm, 0.0)
    term_two = (
        3.0
        / (4 * np.sqrt(5) * np.pi * bandwidth[None, :])
        * bracket
        * logarithm
    )
    term_two = np.where(np.abs(bracket) < 1e-300, 0.0, term_two)
    hilbert_transform = (term_one + term_two).mean(axis=1)
    denominator = (
        np.pi * aspect_ratio * eigenvalues * density
    ) ** 2 + (
        1
        - aspect_ratio
        - np.pi * aspect_ratio * eigenvalues * hilbert_transform
    ) ** 2
    cleaned_eigenvalues = eigenvalues / denominator
    if (
        not np.isfinite(cleaned_eigenvalues).all()
        or (cleaned_eigenvalues <= 0).any()
    ):
        raise RuntimeError("RIE returned non-positive or non-finite eigenvalues.")
    return eigenvectors @ np.diag(cleaned_eigenvalues) @ eigenvectors.T


def load_log_returns(path: Path) -> tuple[pd.Series, list[str], np.ndarray]:
    frame = pd.read_csv(path)
    frame = frame.loc[:, ~frame.columns.duplicated()]
    frame = frame.drop(columns=["Fecha.1"], errors="ignore")
    if "Fecha" not in frame:
        raise ValueError("Input must have a Fecha column and one column per security.")
    frame["Fecha"] = pd.to_datetime(frame["Fecha"], errors="raise")
    frame = frame.sort_values("Fecha").reset_index(drop=True)
    validate_daily_dates(frame["Fecha"])
    tickers = [column for column in frame.columns if column != "Fecha"]
    if not tickers or len(set(tickers)) != len(tickers):
        raise ValueError("Input must contain unique security columns.")
    returns = frame[tickers].apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    if not np.isfinite(returns).all():
        raise ValueError("Input contains missing or non-finite log returns.")
    if (returns < -0.5).any() or (returns > 0.5).any():
        raise ValueError("Unexpected log-return magnitude; verify input units and scale.")
    return frame["Fecha"], tickers, returns


def load_bloomberg_total_returns(
    path: Path,
) -> tuple[pd.Series, list[str], np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        levels_by_ticker: dict[str, pd.Series] = {}
        lower_bound = START_DATE - pd.Timedelta(days=7)
        for ticker, sheet_name in TICKER_SHEET_MAP.items():
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Bloomberg workbook is missing {ticker} ({sheet_name})."
                )
            worksheet = workbook[sheet_name]
            observations = []
            for row in worksheet.iter_rows(
                min_row=7, min_col=1, max_col=16, values_only=True
            ):
                day, level = row[0], row[15]
                if not isinstance(day, (pd.Timestamp,)):
                    if not hasattr(day, "year"):
                        continue
                date = pd.Timestamp(day)
                if lower_bound <= date <= END_DATE and isinstance(
                    level, (int, float)
                ) and np.isfinite(level) and level > 0:
                    observations.append((date, float(level)))
            if not observations:
                raise ValueError(f"No gross total-return observations for {ticker}.")
            series = pd.Series(
                [value for _, value in observations],
                index=pd.DatetimeIndex([date for date, _ in observations]),
                name=ticker,
            )
            if series.index.has_duplicates:
                raise ValueError(f"Duplicate dates in Bloomberg series for {ticker}.")
            levels_by_ticker[ticker] = series.sort_index()
    finally:
        workbook.close()

    levels = pd.concat(levels_by_ticker.values(), axis=1, join="inner").sort_index()
    if levels.isna().any().any():
        raise ValueError("The selected Bloomberg total-return panel has missing values.")
    log_returns = np.log(levels / levels.shift(1)).iloc[1:]
    log_returns = log_returns.loc[
        (log_returns.index >= START_DATE) & (log_returns.index <= END_DATE)
    ]
    if log_returns.empty:
        raise ValueError("No daily log returns remain in the configured date range.")
    dates = pd.Series(log_returns.index, name="Fecha")
    validate_daily_dates(dates)
    return dates, list(log_returns.columns), log_returns.to_numpy(dtype=float)


def load_input(path: Path) -> tuple[pd.Series, list[str], np.ndarray]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_bloomberg_total_returns(path)
    return load_log_returns(path)


def estimate_covariances(training_log_returns: np.ndarray) -> dict[str, np.ndarray]:
    centered = training_log_returns - training_log_returns.mean(axis=0, keepdims=True)
    assets, observations = centered.shape[1], len(centered)
    sample = np.cov(centered, rowvar=False, ddof=0)
    return {
        "Sample": sample,
        "Eigenvalue clipping": eigenvalue_clipping(sample, assets, observations),
        "Ledoit-Wolf": LedoitWolf(assume_centered=True).fit(centered).covariance_,
        "RIE": rie_estimator(sample, assets, observations),
    }


def long_only_weights(covariance: np.ndarray) -> np.ndarray:
    assets = covariance.shape[0]
    initial = np.full(assets, 1 / assets)
    result = minimize(
        lambda weights: weights @ covariance @ weights,
        initial,
        jac=lambda weights: 2 * covariance @ weights,
        method="SLSQP",
        bounds=[(0.0, 1.0)] * assets,
        constraints=[
            {
                "type": "eq",
                "fun": lambda weights: weights.sum() - 1.0,
                "jac": lambda weights: np.ones(assets),
            }
        ],
        options={"maxiter": 500, "ftol": 1e-14},
    )
    if not result.success:
        raise RuntimeError(f"Long-only GMV optimization failed: {result.message}")
    weights = np.clip(result.x, 0.0, None)
    if weights.sum() <= 0:
        raise RuntimeError("Long-only GMV optimizer returned a zero portfolio.")
    return weights / weights.sum()


def breadth(weights: np.ndarray) -> tuple[float, float]:
    effective_number = 1.0 / np.sum(weights**2)
    absolute_weights = np.abs(weights)
    probabilities = absolute_weights / absolute_weights.sum()
    positive = probabilities > 0
    shannon_breadth = np.exp(
        -np.sum(probabilities[positive] * np.log(probabilities[positive]))
    )
    return float(effective_number), float(shannon_breadth)


def summary_row(
    estimator: str,
    constraint: str,
    weights: np.ndarray,
    training_covariance: np.ndarray,
    portfolio_returns: np.ndarray,
    turnover: float = np.nan,
) -> dict[str, object]:
    daily_risk_free = ANNUAL_RISK_FREE_RATE / 252
    daily_volatility = portfolio_returns.std(ddof=1)
    annualized_sharpe = (
        np.sqrt(252) * (portfolio_returns.mean() - daily_risk_free) / daily_volatility
    )
    nea, shannon = breadth(weights)
    return {
        "estimator": estimator,
        "constraint": constraint,
        "training_volatility_pct_per_day": 100
        * np.sqrt(weights @ training_covariance @ weights),
        "holdout_annualized_mean_return_pct": 252 * portfolio_returns.mean() * 100,
        "holdout_volatility_pct_per_day": 100 * daily_volatility,
        "holdout_annualized_volatility_pct": 100 * np.sqrt(252) * daily_volatility,
        "holdout_annualized_sharpe_rf_8pct": annualized_sharpe,
        "effective_number_assets": nea,
        "shannon_effective_breadth": shannon,
        "active_long_positions": int(np.sum(weights > 1e-6)),
        "maximum_weight": float(weights.max()),
        "gross_exposure": float(np.abs(weights).sum()),
        "mean_one_way_turnover_per_rebalance": turnover,
    }


def one_way_turnover(
    prior_weights: np.ndarray,
    asset_log_returns: np.ndarray,
    target_weights: np.ndarray,
) -> float:
    relative_prices = np.exp(asset_log_returns.sum(axis=0))
    wealth = float(relative_prices @ prior_weights)
    if not np.isfinite(wealth) or wealth <= 0:
        raise RuntimeError("Portfolio wealth is non-positive at a rebalance date.")
    drifted_weights = relative_prices * prior_weights / wealth
    return float(0.5 * np.abs(target_weights - drifted_weights).sum())


def load_market_proxies(
    path: Path,
    dates: pd.Series,
    training_length: int,
    test_returns: dict[str, np.ndarray],
) -> pd.DataFrame:
    export = pd.read_csv(path, header=5, dtype=str)
    date_columns = [
        column
        for column in export.columns
        if column == "Date" or column.startswith("Date.")
    ]
    level_columns = [
        column
        for column in export.columns
        if column == "PX_LAST" or column.startswith("PX_LAST.")
    ]
    if len(date_columns) < 4 or len(level_columns) < 3:
        raise ValueError("Market-proxy export lacks the expected Bloomberg columns.")

    def series(date_column: str, level_column: str) -> pd.Series:
        parsed_dates = pd.to_datetime(
            export[date_column], format="%d/%m/%Y", errors="coerce"
        )
        values = pd.to_numeric(
            export[level_column].str.replace(",", ".", regex=False), errors="coerce"
        )
        result = pd.Series(values.to_numpy(), index=parsed_dates)
        return result.loc[result.index.notna()].groupby(level=0).last()

    vomxcus_usd = series(date_columns[0], level_columns[0])
    vomxgus_usd = series(date_columns[1], level_columns[1])
    usd_mxn = series(date_columns[3], level_columns[2])
    proxies = {
        "VOMXCUS": vomxcus_usd.mul(usd_mxn),
        "VOMXGUS": vomxgus_usd.mul(usd_mxn),
    }
    train_dates = pd.DatetimeIndex(dates.iloc[:training_length])
    test_dates = pd.DatetimeIndex(dates.iloc[training_length:])
    rows = []
    for proxy_name, levels in proxies.items():
        threshold_sample = levels.reindex(train_dates).dropna()
        if threshold_sample.empty:
            raise ValueError(f"No training-period observations for {proxy_name}.")
        threshold = float(threshold_sample.quantile(0.9))
        matched = levels.reindex(test_dates).notna()
        state = levels.reindex(test_dates).ge(threshold)
        for state_name, selected in (
            ("Below threshold", matched & ~state),
            ("Above threshold", matched & state),
        ):
            selected_indices = np.flatnonzero(selected.to_numpy())
            row: dict[str, object] = {
                "proxy": proxy_name,
                "state": state_name,
                "matched_days": len(selected_indices),
            }
            for estimator in ESTIMATORS:
                returns = test_returns[f"{estimator} Unconstrained"]
                conditional_returns = returns[selected_indices]
                if len(conditional_returns) < 2:
                    raise ValueError(f"Insufficient matched {proxy_name} observations.")
                row[estimator] = 100 * conditional_returns.std(ddof=1)
            rows.append(row)
    return pd.DataFrame(rows)


def nw_standard_error(influence: np.ndarray, lags: int) -> float:
    count = len(influence)
    centered = influence - influence.mean()
    long_run_variance = np.dot(centered, centered) / count
    for lag in range(1, lags + 1):
        bartlett = 1.0 - lag / (lags + 1.0)
        autocovariance = np.dot(centered[lag:], centered[:-lag]) / count
        long_run_variance += 2 * bartlett * autocovariance
    return float(np.sqrt(max(long_run_variance, 0.0) / count))


def sharpe_ratio(returns: np.ndarray) -> float:
    return float(
        np.sqrt(252)
        * (returns.mean() - ANNUAL_RISK_FREE_RATE / 252)
        / returns.std(ddof=1)
    )


def sharpe_influence(returns: np.ndarray) -> np.ndarray:
    mean = returns.mean()
    volatility = returns.std(ddof=1)
    excess_mean = mean - ANNUAL_RISK_FREE_RATE / 252
    centered = returns - mean
    return np.sqrt(252) * (
        centered / volatility
        - excess_mean * (centered**2 - volatility**2) / (2 * volatility**3)
    )


def volatility_influence(returns: np.ndarray) -> np.ndarray:
    mean = returns.mean()
    volatility = returns.std(ddof=1)
    centered = returns - mean
    return 100 * (centered**2 - volatility**2) / (2 * volatility)


def block_bootstrap_intervals(
    first: np.ndarray,
    second: np.ndarray,
    *,
    block_length: int,
    seed: int,
) -> dict[str, float]:
    count = len(first)
    blocks = int(np.ceil(count / block_length))
    rng = np.random.default_rng(seed)
    sharpe_differences = np.empty(BOOTSTRAP_REPLICATIONS)
    volatility_differences = np.empty(BOOTSTRAP_REPLICATIONS)
    for iteration in range(BOOTSTRAP_REPLICATIONS):
        starts = rng.integers(0, count - block_length + 1, size=blocks)
        indices = np.concatenate(
            [np.arange(start, start + block_length) for start in starts]
        )[:count]
        a, b = first[indices], second[indices]
        sharpe_differences[iteration] = sharpe_ratio(a) - sharpe_ratio(b)
        volatility_differences[iteration] = 100 * (
            a.std(ddof=1) - b.std(ddof=1)
        )
    return {
        "sharpe_difference_lower_95": float(
            np.percentile(sharpe_differences, 2.5)
        ),
        "sharpe_difference_upper_95": float(
            np.percentile(sharpe_differences, 97.5)
        ),
        "volatility_difference_pp_lower_95": float(
            np.percentile(volatility_differences, 2.5)
        ),
        "volatility_difference_pp_upper_95": float(
            np.percentile(volatility_differences, 97.5)
        ),
    }


def inference_rows(
    design: str,
    constraint: str,
    first: np.ndarray,
    second: np.ndarray,
    *,
    blocks: tuple[int, ...],
    base_seed: int,
    benchmark_name: str = "Sample",
) -> list[dict[str, object]]:
    sharpe_difference = sharpe_ratio(first) - sharpe_ratio(second)
    volatility_difference_pp = 100 * (
        first.std(ddof=1) - second.std(ddof=1)
    )
    sharpe_se = nw_standard_error(
        sharpe_influence(first) - sharpe_influence(second), 20
    )
    volatility_se = nw_standard_error(
        volatility_influence(first) - volatility_influence(second), 20
    )
    rows = []
    for block_length in blocks:
        intervals = block_bootstrap_intervals(
            first,
            second,
            block_length=block_length,
            seed=base_seed + block_length,
        )
        for outcome, estimate, se, lower, upper in (
            (
                f"RIE_minus_{benchmark_name}_annualized_Sharpe",
                sharpe_difference,
                sharpe_se,
                intervals["sharpe_difference_lower_95"],
                intervals["sharpe_difference_upper_95"],
            ),
            (
                f"RIE_minus_{benchmark_name}_daily_volatility_percentage_points",
                volatility_difference_pp,
                volatility_se,
                intervals["volatility_difference_pp_lower_95"],
                intervals["volatility_difference_pp_upper_95"],
            ),
        ):
            z_score = estimate / se if se else np.nan
            rows.append(
                {
                    "design": design,
                    "constraint": constraint,
                    "outcome": outcome,
                    "estimate": estimate,
                    "block_length_days": block_length,
                    "bootstrap_lower_95": lower,
                    "bootstrap_upper_95": upper,
                    "hac_lags": 20,
                    "hac_standard_error": se,
                    "hac_z": z_score,
                    "hac_two_sided_p": float(2 * norm.sf(abs(z_score))),
                    "bootstrap_resamples": BOOTSTRAP_REPLICATIONS,
                    "bootstrap_seed": base_seed + block_length,
                }
            )
    return rows


def exact_buy_and_hold_returns(
    asset_log_returns: np.ndarray, weights: np.ndarray
) -> np.ndarray:
    wealth = np.exp(np.cumsum(asset_log_returns, axis=0)) @ weights
    if not np.isfinite(wealth).all() or (wealth <= 0).any():
        raise RuntimeError(
            "Portfolio wealth became non-positive or non-finite during a holding period."
        )
    prior_wealth = np.concatenate(([1.0], wealth[:-1]))
    return wealth / prior_wealth - 1.0


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Run aggregate-only portfolio validation on a lawfully accessed "
            "daily log-return panel. No daily data are written."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        required=True,
        help=(
            "Local daily log-return CSV with Fecha, or a Bloomberg workbook "
            "with total-return-index tabs for the configured 37-stock universe."
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "validation-output",
        help="Folder for aggregate CSV summaries; never use a folder tracked with raw data.",
    )
    parser.add_argument(
        "--market-proxies",
        type=Path,
        help="Optional local Bloomberg CSV export for aggregate VOM regime summaries.",
    )
    arguments = parser.parse_args()
    dates, tickers, log_returns = load_input(arguments.input)
    assets, observations = log_returns.shape[1], len(log_returns)

    if observations < ROLLING_LOOKBACK + ROLLING_REBALANCES * ROLLING_STEP:
        raise ValueError(
            "The return panel is too short for the configured walk-forward design."
        )
    test_length = round(observations * TEST_FRACTION)
    train_length = observations - test_length
    training = log_returns[:train_length]
    test_log_returns = log_returns[train_length:]
    test_simple_returns = np.expm1(test_log_returns)
    covariance_estimates = estimate_covariances(training)
    portfolio_returns: dict[str, np.ndarray] = {}
    weights_by_strategy: dict[str, np.ndarray] = {}
    static_rows = []
    static_turnover: dict[str, float] = {}
    static_turnover_series: dict[str, np.ndarray] = {}

    for name, covariance in covariance_estimates.items():
        for constraint, weights in (
            ("Unconstrained", gmv_portfolio(covariance)),
            ("Long-only", long_only_weights(covariance)),
        ):
            strategy = f"{name} {constraint}"
            weights_by_strategy[strategy] = weights
            portfolio_returns[strategy] = test_simple_returns @ weights
            daily_turnovers = np.zeros(len(test_log_returns))
            for index in range(1, len(test_log_returns)):
                daily_turnovers[index] = one_way_turnover(
                    weights,
                    test_log_returns[index - 1 : index],
                    weights,
                )
            static_turnover_series[strategy] = daily_turnovers
            static_turnover[strategy] = float(daily_turnovers.mean())
            static_rows.append(
                summary_row(
                    name,
                    constraint,
                    weights,
                    covariance,
                    portfolio_returns[strategy],
                    static_turnover[strategy],
                )
            )
    equal_weights = np.full(assets, 1 / assets)
    equal_strategy = "Equal-weight 1/N Long-only"
    weights_by_strategy[equal_strategy] = equal_weights
    portfolio_returns[equal_strategy] = test_simple_returns @ equal_weights
    equal_daily_turnovers = np.zeros(len(test_log_returns))
    for index in range(1, len(test_log_returns)):
        equal_daily_turnovers[index] = one_way_turnover(
            equal_weights, test_log_returns[index - 1 : index], equal_weights
        )
    static_turnover_series[equal_strategy] = equal_daily_turnovers
    static_turnover[equal_strategy] = float(equal_daily_turnovers.mean())
    static_rows.append(
        summary_row(
            "Equal-weight 1/N",
            "Long-only",
            equal_weights,
            covariance_estimates["Sample"],
            portfolio_returns[equal_strategy],
            static_turnover[equal_strategy],
        )
    )

    static_comparisons = []
    for constraint in ("Unconstrained", "Long-only"):
        static_comparisons.extend(
            inference_rows(
                "Fixed training weights; daily target-weight rebalancing",
                constraint,
                portfolio_returns[f"RIE {constraint}"],
                portfolio_returns[f"Sample {constraint}"],
                blocks=(10, 20, 42),
                base_seed=BOOTSTRAP_SEED + (100 if constraint == "Long-only" else 0),
            )
        )
    static_comparisons.extend(
        inference_rows(
            "Fixed training weights; daily target-weight rebalancing",
            "Long-only",
            portfolio_returns["RIE Long-only"],
            portfolio_returns[equal_strategy],
            blocks=(10, 20, 42),
            base_seed=BOOTSTRAP_SEED + 300,
            benchmark_name="Equal-weight_1N",
        )
    )

    rolling_start = observations - ROLLING_REBALANCES * ROLLING_STEP
    rolling_returns: dict[str, list[np.ndarray]] = {}
    rolling_weights: dict[str, list[np.ndarray]] = {}
    rolling_turnovers: dict[str, list[float]] = {}
    rolling_turnover_by_window: dict[str, list[float]] = {}
    dates_of_rebalance = []
    for index in range(ROLLING_REBALANCES):
        start = rolling_start + index * ROLLING_STEP
        training_window = log_returns[start - ROLLING_LOOKBACK : start]
        holding_window = log_returns[start : start + ROLLING_STEP]
        estimates = estimate_covariances(training_window)
        dates_of_rebalance.append(dates.iloc[start])
        for name, covariance in estimates.items():
            for constraint, weights in (
                ("Unconstrained", gmv_portfolio(covariance)),
                ("Long-only", long_only_weights(covariance)),
            ):
                strategy = f"{name} {constraint}"
                rolling_returns.setdefault(strategy, []).append(
                    exact_buy_and_hold_returns(holding_window, weights)
                )
                rolling_weights.setdefault(strategy, []).append(weights)
                turnovers = rolling_turnovers.setdefault(strategy, [])
                window_turnovers = rolling_turnover_by_window.setdefault(strategy, [])
                rebalance_turnover = 0.0
                if index > 0:
                    prior_window = log_returns[
                        start - ROLLING_STEP : start
                    ]
                    prior_weights = rolling_weights[strategy][-2]
                    turnovers.append(
                        one_way_turnover(prior_weights, prior_window, weights)
                    )
                    rebalance_turnover = turnovers[-1]
                window_turnovers.append(rebalance_turnover)
        rolling_returns.setdefault("Equal-weight 1/N Long-only", []).append(
            exact_buy_and_hold_returns(holding_window, equal_weights)
        )
        rolling_weights.setdefault("Equal-weight 1/N Long-only", []).append(
            equal_weights
        )
        equal_turnovers = rolling_turnovers.setdefault(
            "Equal-weight 1/N Long-only", []
        )
        equal_window_turnovers = rolling_turnover_by_window.setdefault(
            "Equal-weight 1/N Long-only", []
        )
        equal_rebalance_turnover = 0.0
        if index > 0:
            prior_window = log_returns[start - ROLLING_STEP : start]
            equal_turnovers.append(
                one_way_turnover(equal_weights, prior_window, equal_weights)
            )
            equal_rebalance_turnover = equal_turnovers[-1]
        equal_window_turnovers.append(equal_rebalance_turnover)

    rolling_rows = []
    pooled_rolling_returns = {}
    for strategy, windows in rolling_returns.items():
        returns = np.concatenate(windows)
        pooled_rolling_returns[strategy] = returns
        prior_weights = rolling_weights[strategy]
        mean_turnover = (
            float(np.mean(rolling_turnovers[strategy]))
            if rolling_turnovers[strategy]
            else 0.0
        )
        if strategy.startswith("Equal-weight"):
            estimator, constraint = "Equal-weight 1/N", "Long-only"
        else:
            estimator, constraint = strategy.rsplit(" ", 1)
        weights_for_summary = prior_weights[-1]
        covariance_for_summary = covariance_estimates.get(
            estimator, covariance_estimates["Sample"]
        )
        rolling_rows.append(
            summary_row(
                estimator,
                constraint,
                weights_for_summary,
                covariance_for_summary,
                returns,
                mean_turnover,
            )
        )

    rolling_comparisons = []
    for constraint in ("Unconstrained", "Long-only"):
        rolling_comparisons.extend(
            inference_rows(
                f"Rolling {ROLLING_LOOKBACK}-day lookback; {ROLLING_STEP}-day buy-and-hold periods",
                constraint,
                pooled_rolling_returns[f"RIE {constraint}"],
                pooled_rolling_returns[f"Sample {constraint}"],
                blocks=(10, 21, 42),
                base_seed=BOOTSTRAP_SEED + (200 if constraint == "Long-only" else 150),
            )
        )
    rolling_comparisons.extend(
        inference_rows(
            f"Rolling {ROLLING_LOOKBACK}-day lookback; {ROLLING_STEP}-day buy-and-hold periods",
            "Long-only",
            pooled_rolling_returns["RIE Long-only"],
            pooled_rolling_returns[equal_strategy],
            blocks=(10, 21, 42),
            base_seed=BOOTSTRAP_SEED + 400,
            benchmark_name="Equal-weight_1N",
        )
    )

    transaction_cost_rows = []
    for design, strategies, turnovers, window_length in (
        (
            "Fixed holdout; daily target weights",
            portfolio_returns,
            static_turnover_series,
            1,
        ),
        (
            f"Rolling; {ROLLING_STEP}-day buy-and-hold",
            pooled_rolling_returns,
            rolling_turnover_by_window,
            ROLLING_STEP,
        ),
    ):
        for strategy, gross_returns in strategies.items():
            estimator, constraint = (
                ("Equal-weight 1/N", "Long-only")
                if strategy.startswith("Equal-weight")
                else strategy.rsplit(" ", 1)
            )
            turnover_values = turnovers[strategy]
            if design.startswith("Fixed"):
                aligned_turnover = turnover_values
                reported_turnovers = turnover_values
            else:
                aligned_turnover = np.zeros(len(gross_returns))
                for window_index, turnover in enumerate(turnover_values):
                    aligned_turnover[window_index * window_length] = turnover
                reported_turnovers = turnover_values[1:]
            for cost_bps in (0, 5, 10, 25, 50):
                net_returns = gross_returns - aligned_turnover * cost_bps / 10_000
                transaction_cost_rows.append(
                    {
                        "design": design,
                        "estimator": estimator,
                        "constraint": constraint,
                        "one_way_cost_basis_points": cost_bps,
                        "mean_one_way_turnover_per_rebalance": float(
                            np.mean(reported_turnovers)
                        ),
                        "gross_annualized_mean_return_pct": float(
                            252 * gross_returns.mean() * 100
                        ),
                        "net_annualized_mean_return_pct": float(
                            252 * net_returns.mean() * 100
                        ),
                        "net_annualized_sharpe_rf_8pct": sharpe_ratio(net_returns),
                        "gross_annualized_sharpe_rf_8pct": sharpe_ratio(gross_returns),
                    }
                )

    arguments.output.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(static_rows).to_csv(
        arguments.output / "static_constraint_baselines.csv", index=False
    )
    primary_rows = [
        {
            "estimator": row["estimator"],
            "in_sample_volatility_pct_per_day": row[
                "training_volatility_pct_per_day"
            ],
            "holdout_volatility_pct_per_day": row[
                "holdout_volatility_pct_per_day"
            ],
            "annualized_sharpe": row["holdout_annualized_sharpe_rf_8pct"],
            "effective_number_assets": row["effective_number_assets"],
            "shannon_effective_breadth": row["shannon_effective_breadth"],
            "holdout_annualized_mean_return_pct": row[
                "holdout_annualized_mean_return_pct"
            ],
        }
        for row in static_rows
        if row["constraint"] == "Unconstrained"
    ]
    pd.DataFrame(primary_rows).to_csv(
        arguments.output / "primary_outcomes.csv", index=False
    )
    pd.DataFrame(rolling_rows).to_csv(
        arguments.output / "rolling_walkforward_outcomes.csv", index=False
    )
    pd.DataFrame(transaction_cost_rows).to_csv(
        arguments.output / "transaction_cost_sensitivity.csv", index=False
    )
    pd.DataFrame(static_comparisons + rolling_comparisons).to_csv(
        arguments.output / "inference_reanalysis.csv", index=False
    )
    inference_frame = pd.DataFrame(static_comparisons + rolling_comparisons)
    selected_inference = inference_frame.loc[
        (
            inference_frame["design"]
            == "Fixed training weights; daily target-weight rebalancing"
        )
        & (inference_frame["block_length_days"] == 20)
        | (
            inference_frame["design"]
            == "Rolling 504-day lookback; 21-day buy-and-hold periods"
        )
        & (inference_frame["block_length_days"] == 21)
    ].copy()
    selected_inference["quantity"] = (
        selected_inference["design"].map(
            {
                "Fixed training weights; daily target-weight rebalancing": "Fixed-holdout",
                "Rolling 504-day lookback; 21-day buy-and-hold periods": "Rolling-exploratory",
            }
        )
        + "_"
        + selected_inference["constraint"]
        + "_"
        + selected_inference["outcome"]
    )
    selected_inference["lower_95"] = selected_inference["bootstrap_lower_95"]
    selected_inference["upper_95"] = selected_inference["bootstrap_upper_95"]
    selected_inference["p_value"] = selected_inference["hac_two_sided_p"]
    selected_inference["details"] = (
        selected_inference["design"]
        + "; paired moving-block percentile 95% CI; HAC(20) two-sided p-value; "
        + selected_inference["bootstrap_resamples"].astype(str)
        + " resamples"
    )
    selected_inference[
        ["quantity", "estimate", "lower_95", "upper_95", "p_value", "details"]
    ].to_csv(arguments.output / "inference_summary.csv", index=False)
    if arguments.market_proxies:
        load_market_proxies(
            arguments.market_proxies, dates, train_length, portfolio_returns
        ).to_csv(arguments.output / "market_regimes.csv", index=False)
    pd.DataFrame(
        {
            "rebalance_date": pd.to_datetime(dates_of_rebalance).strftime("%Y-%m-%d"),
            "lookback_days": ROLLING_LOOKBACK,
            "holding_days": ROLLING_STEP,
        }
    ).to_csv(arguments.output / "rolling_design.csv", index=False)
    pd.DataFrame(
        [
            {
                "asset_count": assets,
                "observation_count": observations,
                "first_return_date": dates.iloc[0].date(),
                "last_return_date": dates.iloc[-1].date(),
                "train_observations": train_length,
                "test_observations": test_length,
                "daily_gap_count_over_7_calendar_days": int(
                    (dates.sort_values().diff().dt.days > 7).sum()
                ),
                "return_source": (
                    "Bloomberg TOT_RETURN_INDEX_GROSS_DVDS"
                    if arguments.input.suffix.lower() in {".xlsx", ".xlsm"}
                    else "Input daily log-return CSV"
                ),
                "selected_universe": "37 mapped equities from the prior 38-stock list; FUNO11 absent from workbook",
            }
        ]
    ).to_csv(arguments.output / "panel_audit.csv", index=False)
    pd.DataFrame(
        {"economatica_ticker": tickers, "bloomberg_sheet": [
            TICKER_SHEET_MAP.get(ticker, "Input CSV")
            for ticker in tickers
        ]}
    ).to_csv(arguments.output / "universe_mapping.csv", index=False)
    print(f"Validated {assets} assets and {observations} daily log-return observations.")
    print(
        f"Fixed split: {train_length}/{test_length}; "
        f"{dates.iloc[train_length].date()} to {dates.iloc[-1].date()}. "
        f"Rolling periods: {dates.iloc[rolling_start].date()} "
        f"to {dates.iloc[-1].date()} ({ROLLING_REBALANCES} rebalances)."
    )
    print("Fixed-holdout outcomes use exact simple asset returns at target daily weights.")
    print("Rolling outcomes hold shares fixed inside each 21-day block and rebalance monthly.")
    print("Only aggregate summaries were written; no daily inputs or returns were saved.")
    print(f"Aggregate outputs: {arguments.output.resolve()}")


if __name__ == "__main__":
    main()
